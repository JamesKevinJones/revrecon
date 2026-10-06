"""Build the Excel audit workpaper from the generated CSVs.

Re-performs the reconciliation in Python (independently of the SQL views) and
writes workpapers/RevRecon_Reconciliation_Workpaper.xlsx:

  Cover & Sign-off   objective, procedures, tickmark legend, preparer/reviewer
  Summary            exception counts and amounts as live SUMIFS formulas
  Recon Detail       every contract-period: expected vs billed
  Exceptions         the non-reconciled rows, with follow-up columns to fill in

    python scripts/build_workpaper.py
"""

from __future__ import annotations

import argparse
from decimal import Decimal
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from billing_rules import classify, expected_charge, sla_credit_pct

ROOT = Path(__file__).resolve().parent.parent
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")   # cells the analyst fills in
MONEY = '#,##0.00;[Red]-#,##0.00'

EXCEPTION_TYPES = ["UNBILLED_USAGE", "TIER_PRICING_VARIANCE", "MISSED_SLA_CREDIT", "RECONCILED"]

DETAIL_COLUMNS = [
    ("Customer ID", 12, None), ("Company", 32, None), ("Tier", 11, None), ("Contract ID", 12, None),
    ("Billing Period", 14, "yyyy-mm"), ("Events", 9, "#,##0"), ("Units", 10, "#,##0"),
    ("Billable Units", 13, "#,##0"), ("Availability %", 13, "0.000"), ("SLA Target %", 12, "0.00"),
    ("Credit %", 9, "0%"), ("Gross Expected", 15, MONEY), ("SLA Credit", 12, MONEY),
    ("Net Expected", 14, MONEY), ("Invoice ID", 11, None), ("Billed", 13, MONEY),
    ("Leakage", 12, MONEY), ("Exception Type", 24, None), ("Tickmark", 10, None),
]


def reperform(data: Path) -> pd.DataFrame:
    customers = pd.read_csv(data / "dim_customer.csv")
    contracts = pd.read_csv(data / "dim_contract.csv", dtype={"base_rate": str, "tier_discount_rate": str, "sla_uptime_target": str})
    events = pd.read_csv(data / "fact_usage_events.csv", parse_dates=["timestamp"])
    invoices = pd.read_csv(data / "fact_invoices.csv", dtype={"billed_amount": str}, parse_dates=["billing_period"])

    events["billing_period"] = events["timestamp"].dt.to_period("M").dt.to_timestamp()
    usage = (events.groupby(["customer_id", "contract_id", "billing_period"])
             .agg(events=("event_id", "size"), units=("units_consumed", "sum"), errors=("is_error", "sum"))
             .reset_index())
    df = (usage.merge(contracts, on=["contract_id", "customer_id"])
               .merge(customers[["customer_id", "company_name", "tier"]], on="customer_id")
               .merge(invoices[["contract_id", "billing_period", "invoice_id", "billed_amount"]],
                      on=["contract_id", "billing_period"], how="left"))

    rows = []
    for r in df.itertuples(index=False):
        target = Decimal(r.sla_uptime_target)
        credit_pct = sla_credit_pct(int(r.events), int(r.errors), target)
        gross, credit, net = expected_charge(int(r.units), int(r.committed_volume),
                                             Decimal(r.base_rate), Decimal(r.tier_discount_rate), credit_pct)
        billed = None if pd.isna(r.billed_amount) else Decimal(r.billed_amount)
        rows.append([
            r.customer_id, r.company_name, r.tier, r.contract_id, r.billing_period.to_pydatetime(),
            int(r.events), int(r.units), max(int(r.units), int(r.committed_volume)),
            round(100 * (r.events - r.errors) / r.events, 3), float(target), float(credit_pct),
            float(gross), float(credit), float(net),
            None if pd.isna(r.invoice_id) else int(r.invoice_id),
            None if billed is None else float(billed),
            float(net - (billed or 0)), classify(billed, gross, credit, net), None,
        ])
    detail = pd.DataFrame(rows, columns=[c[0] for c in DETAIL_COLUMNS])
    return detail.sort_values(["Billing Period", "Contract ID"]).reset_index(drop=True)


def write_table(ws: Worksheet, frame: pd.DataFrame, columns: list[tuple], start_row: int = 1) -> None:
    for col, (name, width, _) in enumerate(columns, start=1):
        cell = ws.cell(row=start_row, column=col, value=name)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    for r, values in enumerate(frame.itertuples(index=False), start=start_row + 1):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=None if pd.isna(value) else value)
            if isinstance(value, str) and value.startswith("="):
                # openpyxl stores a leading "=" as a live formula. Source text
                # (e.g. a company name) must stay text, or it is formula injection.
                cell.data_type = "s"
            if columns[c - 1][2]:
                cell.number_format = columns[c - 1][2]
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(columns))}{start_row + len(frame)}"


def cover_sheet(ws: Worksheet, detail: pd.DataFrame) -> None:
    first, last = detail["Billing Period"].min(), detail["Billing Period"].max()
    lines = [
        ("RevRecon: Revenue Leakage & SLA Billing Reconciliation", Font(bold=True, size=14)),
        (f"Period under review: {first:%b %Y} to {last:%b %Y}", None),
        ("", None),
        ("Objective", Font(bold=True)),
        ("Confirm every contract-period with metered usage was invoiced at the contracted tier price, "
         "net of any SLA credit owed (BRD BR-01 to BR-05).", None),
        ("", None),
        ("Procedures", Font(bold=True)),
        ("1. Re-perform expected billing from usage events and contract terms (Recon Detail).", None),
        ("2. Agree each expected amount to the invoice register; investigate differences over 0.01.", None),
        ("3. Classify each exception by root cause and assign an owner (Exceptions).", None),
        ("4. Agree Summary totals to vw_billing_reconciliation in MySQL and to the Power BI Executive Summary.", None),
        ("", None),
        ("Tickmark legend", Font(bold=True)),
        ("R   Re-performed: recalculated from source data without exception", None),
        ("A   Agreed to invoice register", None),
        ("E   Exception: see Exceptions sheet", None),
        ("S   Agreed to SQL view output", None),
        ("", None),
        ("Sign-off", Font(bold=True)),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        if font:
            cell.font = font
    row = len(lines) + 1
    for label in ("Prepared by", "Date prepared", "Reviewed by", "Date reviewed"):
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2).fill = INPUT_FILL
        row += 1
    ws.column_dimensions["A"].width = 100
    ws.column_dimensions["B"].width = 24


def summary_sheet(ws: Worksheet, detail_rows: int) -> None:
    last = detail_rows + 1
    leakage = f"'Recon Detail'!$Q$2:$Q${last}"
    kind = f"'Recon Detail'!$R$2:$R${last}"
    net = f"'Recon Detail'!$N$2:$N${last}"
    billed = f"'Recon Detail'!$P$2:$P${last}"

    columns = [("Exception Type", 26, None), ("Count", 10, "#,##0"), ("Net Leakage", 15, MONEY),
               ("Under-billed", 15, MONEY), ("Over-billed", 15, MONEY)]
    write_table(ws, pd.DataFrame(columns=[c[0] for c in columns]), columns)
    ws.freeze_panes = None
    ws.auto_filter.ref = None
    for i, name in enumerate(EXCEPTION_TYPES, start=2):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=f'=COUNTIFS({kind},A{i})')
        ws.cell(row=i, column=3, value=f'=SUMIFS({leakage},{kind},A{i})')
        ws.cell(row=i, column=4, value=f'=SUMIFS({leakage},{kind},A{i},{leakage},">0")')
        ws.cell(row=i, column=5, value=f'=-SUMIFS({leakage},{kind},A{i},{leakage},"<0")')
    total = len(EXCEPTION_TYPES) + 2
    ws.cell(row=total, column=1, value="Total").font = Font(bold=True)
    for col in "BCDE":
        ws[f"{col}{total}"] = f"=SUM({col}2:{col}{total - 1})"
        ws[f"{col}{total}"].font = Font(bold=True)

    kpis = [
        ("Total Expected Revenue", f"=SUM({net})", MONEY),
        ("Total Billed Revenue", f"=SUM({billed})", MONEY),
        ("Net Leakage Variance", "=B11-B12", MONEY),
        ("% Leakage Exposure", "=IF(B11=0,0,B13/B11)", "0.00%"),
        ("Gross Exposure (under + over)", f"=D{total}+E{total}", MONEY),
    ]
    for i, (label, formula, fmt) in enumerate(kpis, start=11):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=i, column=2, value=formula)
        cell.number_format = fmt
    for row in range(2, total + 1):
        for col in (3, 4, 5):
            ws.cell(row=row, column=col).number_format = MONEY


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--data", type=Path, default=ROOT / "data")
    parser.add_argument("--out", type=Path, default=ROOT / "workpapers" / "RevRecon_Reconciliation_Workpaper.xlsx")
    args = parser.parse_args()

    detail = reperform(args.data)

    # Re-performance must land on the generator's ground truth exactly.
    manifest = pd.read_csv(args.data / "anomaly_manifest.csv", parse_dates=["billing_period"])
    found = detail.loc[detail["Exception Type"] != "RECONCILED", ["Contract ID", "Billing Period", "Exception Type"]]
    found = set(map(tuple, found.astype(str).to_numpy()))
    seeded = set(map(tuple, manifest[["contract_id", "billing_period", "anomaly_type"]].astype(str).to_numpy()))
    assert found == seeded, f"re-performance disagrees with seeded anomalies: {len(found ^ seeded)} rows differ"

    wb = Workbook()
    cover_sheet(wb.active, detail)
    wb.active.title = "Cover & Sign-off"
    summary_sheet(wb.create_sheet("Summary"), len(detail))
    write_table(wb.create_sheet("Recon Detail"), detail, DETAIL_COLUMNS)

    exceptions = detail[detail["Exception Type"] != "RECONCILED"].drop(columns=["Tickmark"]).copy()
    for col in ("Root Cause Confirmed", "Owner", "Resolution", "Status"):
        exceptions[col] = None
    exc_columns = DETAIL_COLUMNS[:-1] + [("Root Cause Confirmed", 18, None), ("Owner", 16, None),
                                         ("Resolution", 30, None), ("Status", 12, None)]
    ws = wb.create_sheet("Exceptions")
    write_table(ws, exceptions, exc_columns)
    for row in range(2, len(exceptions) + 2):
        for col in range(len(DETAIL_COLUMNS), len(exc_columns) + 1):
            ws.cell(row=row, column=col).fill = INPUT_FILL

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"Wrote {args.out}")
    print(f"  {len(detail)} contract-periods, {len(exceptions)} exceptions "
          f"(re-performance matches all {len(seeded)} seeded anomalies)")


if __name__ == "__main__":
    main()
